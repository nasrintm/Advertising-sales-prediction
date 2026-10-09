"""
Builds Assignment2_Report.docx (A4) from results.json, model_comparison.csv and plots/.
Run after train_model.py:   python build_report.py
All numbers in the text come from the result files, so they always match the code.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
R = json.loads((BASE / "results.json").read_text())
PLOTS = BASE / "plots"

BASE_M, POLY_M, RF_M = ("Multiple Linear Regression", "Polynomial Linear Regression (deg 2)", "Random Forest")
T, TR, CV = R["test"], R["train"], R["cv"]
D, P, L = R["describe"], R["preprocessing"], R["original_labels"]
coef = R["baseline_coef"]
final = R["final_model"]
assert final == POLY_M, "report text assumes the polynomial model won - re-read results.json"

rmse_cut = (1 - T[POLY_M]["RMSE"] / T[BASE_M]["RMSE"]) * 100
mae_cut = (1 - T[POLY_M]["MAE"] / T[BASE_M]["MAE"]) * 100
radio_vs_tv = coef["Radio"] / coef["TV"]
rf_gap = TR[RF_M]["R2"] - T[RF_M]["R2"]
poly_gap = TR[POLY_M]["R2"] - T[POLY_M]["R2"]
mean_sales = D["Sales"]["mean"]
cs = R["correlation_with_sales"]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)          # A4
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
for lvl, size in ((1, 13.5), (2, 11.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x8C)
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(9), Pt(3)

# page number in the footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
    run._r.append(el)
run.font.size = Pt(9)


def para(text="", bold=False, italic=False, align=None, size=None, color=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def rich(parts, style=None):
    """parts = [('text', bold?), ...]"""
    p = doc.add_paragraph(style=style)
    for text, b in parts:
        p.add_run(text).bold = b
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(rows, widths, bold_rows=(), size=9.5, highlight_row=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(val))
            r.font.size = Pt(size)
            r.bold = i == 0 or i in bold_rows
            if j > 0 and i > 0 and str(val).replace(".", "", 1).replace("-", "", 1).isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                shade(c, "DCE6F4")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            elif highlight_row == i:
                shade(c, "FFF2E8")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(name, caption, width):
    doc.add_picture(str(PLOTS / name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(1)
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(caption, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, color=(0x4B, 0x55, 0x63), after=6)


def figure_pair(items, width):
    """items = [(file, caption), (file, caption)] side by side."""
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (name, cap) in enumerate(items):
        c = t.cell(0, j); c.text = ""
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run().add_picture(str(PLOTS / name), width=Inches(width))
        c2 = t.cell(1, j); c2.text = ""
        p = c2.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(cap); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ================================================================= title
para("Machine Learning Application - Assignment 2", bold=True, size=19, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Advertising Sales Prediction using Linear Regression", size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
     color=(0x1F, 0x4E, 0x8C), after=2)
para("Name: ______________________     Roll No.: ______________     Class / Section: ______________     "
     "Date: ______________", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, after=6)

# ================================================================= 1
doc.add_heading("1. Problem Statement", 1)
doc.add_paragraph(
    "A company spends its advertising money on three channels: TV, radio and newspaper. Management wants to know "
    "how much each channel contributes to sales and to predict the sales that a planned budget will produce, so "
    "that money can be moved to the channels that work best. This is a supervised regression problem, because "
    "the value to predict (sales) is a continuous number. The aim of this project is to build, test and compare "
    "linear regression models that predict sales from the three advertising budgets, and to package the best "
    "model in a small desktop application.")

# ================================================================= 2
doc.add_heading("2. Dataset Description", 1)
doc.add_paragraph(
    f"The file Advertising_Budget_and_Sales.csv has {P['rows_raw']} rows (one per market) and four numeric "
    "columns, plus an unnamed index column that is dropped. It is the well-known Advertising dataset used in "
    "'An Introduction to Statistical Learning'. There are three input features and one target:")
table([
    ["Column in the file", "Short name", "Role", "Min", "Mean", "Max"],
    [L["TV"], "TV", "Feature", D["TV"]["min"], f"{D['TV']['mean']:.1f}", D["TV"]["max"]],
    [L["Radio"], "Radio", "Feature", D["Radio"]["min"], f"{D['Radio']['mean']:.1f}", D["Radio"]["max"]],
    [L["Newspaper"], "Newspaper", "Feature", D["Newspaper"]["min"], f"{D['Newspaper']['mean']:.1f}", D["Newspaper"]["max"]],
    [L["Sales"], "Sales", "Target", D["Sales"]["min"], f"{D['Sales']['mean']:.1f}", D["Sales"]["max"]],
], widths=[2.1, 1.0, 0.8, 0.7, 0.7, 0.7])
para("Units: the file labels all four columns in dollars ($). In the original source the budgets are in "
     "thousands of dollars and sales in thousands of units. This report uses the units exactly as labelled "
     "in the file.", italic=True, size=9.5, color=(0x4B, 0x55, 0x63))
doc.add_paragraph(
    "Because the target is a number, there are no classes, so the 'class counts' of a classification problem do "
    "not apply. Instead the target distribution is shown in Figure 1: sales range from "
    f"{D['Sales']['min']} to {D['Sales']['max']} with a mean of {mean_sales:.2f}.")
figure("01_distributions.png", "Figure 1: Distribution of the three budgets and of sales.", 6.5)

# ================================================================= 3
doc.add_heading("3. Data Preprocessing", 1)
bullets([
    ("Missing values and duplicates: ", f"{P['missing_values']} missing values and {P['duplicate_rows']} duplicate "
     f"rows were found, so all {P['rows_clean']} rows were kept."),
    ("Text columns: ", f"{P['text_columns']} - every column is numeric, so no encoding was needed. The unnamed "
     "index column was dropped and the long headers were shortened to TV, Radio, Newspaper and Sales."),
    ("Outliers: ", "checked with the 1.5 x IQR rule: " + ", ".join(f"{k} {v}" for k, v in P["iqr_outliers"].items()) +
     ". The two newspaper points are large but realistic budgets, so they were kept."),
    ("Scaling: ", "StandardScaler is placed inside a scikit-learn Pipeline, so the mean and standard deviation "
     "are learned from the training data only (no data leakage)."),
    ("Split: ", f"80% training ({R['n_train']} rows) and 20% testing ({R['n_test']} rows) with random_state = 42. "
     "stratify= is used only for classification, so it is not used for a numeric target. The test set was "
     "kept untouched until the final evaluation."),
])
figure_pair([("02_correlation_heatmap.png", "Figure 2: Correlation matrix."),
             ("07_tv_radio_interaction.png", "Figure 3: TV effect for low, medium and high radio budgets.")], 3.15)
doc.add_paragraph(
    f"TV has the strongest correlation with sales (r = {cs['TV']:.2f}), followed by radio (r = {cs['Radio']:.2f}) "
    f"and newspaper (r = {cs['Newspaper']:.2f}). Radio and newspaper are themselves correlated "
    f"(r = {R['radio_newspaper_corr']:.2f}), which explains why newspaper looks useful on its own but adds little "
    "once radio is known. Figure 3 shows an important pattern: the sales gained from extra TV spending are larger "
    "when the radio budget is high (the green line is steeper). The two channels reinforce each other.")

# ================================================================= 4
doc.add_heading("4. Algorithm Selection and Why", 1)
doc.add_paragraph("Three algorithms were compared. The topic of the project is linear regression, so two of "
                  "the three are linear regression models:")
table([
    ["Algorithm", "Why it was chosen"],
    [BASE_M, "Baseline. Sales = b0 + b1 TV + b2 Radio + b3 Newspaper. Simple and easy to explain: each coefficient is the "
             "change in sales for one extra unit of that budget."],
    [POLY_M, "Still linear regression (it is linear in its coefficients) but with squared terms and products such as "
             "TV x Radio added, so it can capture the way the channels reinforce each other (Figure 3)."],
    [RF_M, "Non-linear benchmark that needs no assumption about the shape of the relationship. It shows whether "
           "a more flexible model would do better."],
], widths=[2.2, 4.6])

# ================================================================= 5
doc.add_heading("5. Training and Testing", 1)
doc.add_paragraph(
    "Each algorithm was built as a Pipeline and trained with .fit() on the training set. To choose the final "
    "model, 5-fold cross-validation was run on the training set only, and the model with the lowest "
    "cross-validated RMSE among the linear regression models was selected. The test set was then used once, "
    "with .predict(), to measure performance on unseen markets. Choosing by cross-validation instead of by "
    "the test score keeps the test result an honest estimate.")

# ================================================================= 6
doc.add_heading("6. Evaluation Metrics", 1)
doc.add_paragraph(
    "Accuracy, precision, recall, F1, ROC-AUC and the confusion matrix are classification metrics and cannot "
    "be used for a numeric target. The regression equivalents are: MAE (average size of the error), MSE and RMSE "
    "(squared error; RMSE is in the same unit as sales), R² (share of the variation in sales that the model "
    "explains) and adjusted R² (R² corrected for the number of predictors).")
rows = [["Model", "CV RMSE", "CV R²", "Test MAE", "Test MSE", "Test RMSE", "Test R²", "Adj. R²"]]
short = {BASE_M: "Multiple Linear Regression", POLY_M: "Polynomial Linear Reg. (deg 2)", RF_M: "Random Forest"}
for m in (BASE_M, POLY_M, RF_M):
    rows.append([short[m], f"{CV[m]['cv_rmse_mean']:.3f}", f"{CV[m]['cv_r2_mean']:.3f}", f"{T[m]['MAE']:.3f}",
                 f"{T[m]['MSE']:.3f}", f"{T[m]['RMSE']:.3f}", f"{T[m]['R2']:.3f}", f"{T[m]['Adj_R2']:.3f}"])
table(rows, widths=[2.0, 0.72, 0.66, 0.72, 0.72, 0.78, 0.7, 0.7], size=9, highlight_row=2)
para("Table 1: Cross-validation (training set, 5 folds) and test-set results. The shaded row is the final model.",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x4B, 0x55, 0x63))

# ================================================================= 7
doc.add_heading("7. Results and Graphs", 1)
figure("06_model_comparison.png", "Figure 4: Model comparison - cross-validated RMSE (left) and test RMSE (right); lower is better.", 6.3)
figure("04_actual_vs_predicted.png", "Figure 5: Actual vs predicted sales on the test set; points on the dashed line are perfect predictions.", 6.2)
figure("05_residuals.png", "Figure 6: Residuals (actual - predicted) of the baseline (top) and the final model (bottom).", 5.3)

# ================================================================= 8
doc.add_heading("8. Analysis of Results", 1)
bullets([
    ("Which model won: ", f"{POLY_M}. It had the lowest cross-validated RMSE ({CV[POLY_M]['cv_rmse_mean']:.3f} vs "
     f"{CV[RF_M]['cv_rmse_mean']:.3f} for Random Forest and {CV[BASE_M]['cv_rmse_mean']:.3f} for the baseline) and the best "
     f"test result: R² = {T[POLY_M]['R2']:.3f}, RMSE = {T[POLY_M]['RMSE']:.3f}, MAE = {T[POLY_M]['MAE']:.3f}. "
     f"Compared with the plain multiple regression it cuts the test RMSE by {rmse_cut:.0f}% and the MAE by {mae_cut:.0f}%."),
    ("What the errors mean: ", f"the final model is typically wrong by about {T[POLY_M]['MAE']:.2f} sales units "
     f"({T[POLY_M]['MAE'] / mean_sales * 100:.0f}% of the average sales of {mean_sales:.1f}), and the larger misses "
     f"are about {T[POLY_M]['RMSE']:.2f} (RMSE). The baseline is off by about {T[BASE_M]['MAE']:.2f} on average. "
     "In Figure 6 the baseline residuals form a U-shape: it under-predicts at very low and very high sales and "
     "over-predicts in the middle, which means it misses a pattern. The final model's residuals are scattered "
     "evenly around zero with no pattern."),
    ("Which metric matters most: ", "for budgeting, RMSE and MAE matter most because they say how many sales "
     "units a forecast can be wrong by, in the unit of sales; RMSE punishes large misses, which are the costly "
     "ones when a budget is planned. R² is useful for comparing models but does not say how large the error is."),
    ("Random Forest: ", f"nearly as accurate on the test set (R² = {T[RF_M]['R2']:.3f}) but it fits the training "
     f"data much more closely (training R² {TR[RF_M]['R2']:.3f} vs test {T[RF_M]['R2']:.3f}, a gap of {rf_gap:.3f}) "
     f"than the polynomial model (training {TR[POLY_M]['R2']:.3f} vs test {T[POLY_M]['R2']:.3f}, no sign of "
     "overfitting), and it cannot be written as a simple equation. The polynomial linear regression is "
     "therefore the better choice here."),
    ("Business insight (baseline coefficients): ", f"Sales = {R['baseline_intercept']:.3f} + {coef['TV']:.4f} x TV + "
     f"{coef['Radio']:.4f} x Radio + {coef['Newspaper']:.4f} x Newspaper. One extra unit of radio budget is "
     f"associated with about {radio_vs_tv:.1f} times the sales gain of one extra unit of TV budget, while newspaper "
     "has almost no effect once TV and radio are known."),
])
doc.add_heading("Limitations", 2)
bullets([
    f"Small data: with only {R['n_train'] + R['n_test']} markets and a {R['n_test']}-row test set, the metrics "
    f"carry some uncertainty (the baseline's cross-validated R² varies by ±{CV[BASE_M]['cv_r2_std']:.2f} between folds).",
    "Predictions outside the budget ranges in the data are unreliable, and squared terms can behave badly there. "
    "The desktop application warns the user when a value is out of range.",
    "The model shows association, not proof of cause. Price, season, region and competitors are not in the data.",
    "The polynomial model's individual coefficients are harder to interpret than those of the baseline.",
])

# ================================================================= 9
doc.add_heading("9. Mapping to the Course Outcome", 1)
rich([("CO5 - Solve real life problems using appropriate machine learning models and evaluate the performance "
       "measures. ", True),
      ("This project addresses a real marketing problem (predicting sales from advertising spend), compares three "
       "suitable models and selects one using cross-validation, and evaluates it with MAE, MSE, RMSE, R², "
       "adjusted R² and residual analysis.", False)])

# ================================================================= 10
doc.add_heading("10. Desktop Application", 1)
doc.add_paragraph(
    "The final model is packaged as a desktop application (desktop_app.py, built with Tkinter). The user types or "
    "slides the three budgets and sees the predicted sales, the typical error, the prediction of the other models "
    "for comparison and a warning when an input is outside the training range. Other tabs show the dataset, the "
    "model results and all graphs, and the prediction history can be exported to CSV. To run the project: "
    "pip install pandas numpy matplotlib seaborn scikit-learn joblib python-docx; python train_model.py; "
    "python desktop_app.py.")
t = doc.add_table(rows=1, cols=1)
t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
c = t.cell(0, 0); c.text = ""
p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n[ Insert a screenshot of the running application here ]\n"); r.italic = True
r.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
para("Figure 7: Screenshot of the desktop application (Predict tab).", italic=True, size=9,
     align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x4B, 0x55, 0x63))

doc.save(BASE / "Assignment2_Report.docx")
print("Saved Assignment2_Report.docx")
